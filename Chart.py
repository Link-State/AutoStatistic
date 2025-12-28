
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment
from openpyxl.styles import Font as CellFont
from openpyxl.styles import Border, Side
from openpyxl.cell.cell import Cell
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from Stat import Stat
from Table import Table, FundamentalTable, CorrelationTable

# https://ybworld.tistory.com/147

class Chart() :
    def __init__(self, stat:Stat, table:Table, location:str) :
        self.stat = stat
        self.table = table
        self.location = location
        self.areaTuple = table.area
        self.sheet:Worksheet = self.table.connectedSheet
        self.titleArea = None
        self.dataArea = None
        self.chart = None

        sheet = table.connectedSheet
        titleStart:Cell = sheet[self.areaTuple[0]]
        titleEnd:Cell = titleStart.offset(0, table.width-1)
        self.titleArea = titleStart.coordinate + ":" + titleEnd.coordinate

        return
    
    def inflate(self, sh=None) :
        # 테이블 제목 셀 병합
        self.sheet.merge_cells(self.titleArea)
        range = self.sheet[self.titleArea]

        # 제목 셀 가운데 정렬 및 굵게
        title_font = CellFont(name='맑은 고딕', bold=True)
        title_alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
        title_border = Border(left=Side(border_style='thin'),
                                     top=Side(border_style='thin'),
                                     right=Side(border_style='thin'),
                                     bottom=Side(border_style='thin'))
        for rows in range :
            for cell in rows :
                cell.font = title_font
                cell.alignment = title_alignment
                cell.border = title_border
        
        self.table.fitMergedCellSize(self.titleArea)

        return

class FundamentalChart(Chart) :
    def __init__(self, stat:Stat, table:FundamentalTable, location:str) :
        super().__init__(stat=stat, table=table, location=location)
        self.categoryArea = None

        startCell:Cell = self.sheet[self.areaTuple[0]]
        endCell:Cell = self.sheet[self.areaTuple[1]]

        catStart = startCell.offset(1, 0)
        catEnd = catStart.offset(endCell.row - catStart.row, 0)
        self.categoryArea = catStart.coordinate + ":" + catEnd.coordinate

        dataStart = startCell.offset(1, 1)
        dataEnd:Cell = self.sheet[endCell.coordinate]
        self.dataArea = dataStart.coordinate + ":" + dataEnd.coordinate

        return
    
    def inflate(self, sh=None) :
        sheet:Worksheet = self.sheet
        if isinstance(sh, Worksheet) :
            sheet = sh
        else :
            super().inflate()
        
        # 차트 제목 글꼴 설정
        title_font = Font(typeface="맑은 고딕")
        title_character_properties = CharacterProperties(latin=title_font, sz=1800, b=True)
        title_paragraph_properties = ParagraphProperties(defRPr=title_character_properties)

        # 차트 준비
        self.chart = PieChart()
        self.chart.roundedCorners = False
        self.chart.title = self.table.title
        self.chart.title.tx.rich.p[0].pPr = title_paragraph_properties
        self.chart.title.overlay = False
        self.chart.style = 2
        self.chart.legend = None

        # 참조데이터 준비        
        dataStart:Cell = self.sheet[self.dataArea][0][0]
        dataEnd:Cell = self.sheet[self.dataArea][-1][-1]

        catStart:Cell = self.sheet[self.categoryArea][0][0]
        catEnd:Cell = self.sheet[self.categoryArea][-1][-1]

        data = Reference(self.sheet, min_col=dataStart.column, min_row=dataStart.row-1, max_row=dataEnd.row)
        cats = Reference(self.sheet, min_col=catStart.column, min_row=catStart.row, max_row=catEnd.row)

        # 차트 데이터 및 카테고리 추가
        self.chart.add_data(data, titles_from_data=True)
        self.chart.set_categories(cats)
        self.chart.width = 12.8524
        self.chart.height = 10.1854

        # 차트 데이터라벨 글꼴 설정
        fontsize = 700
        fontbold = False
        dataAmount = dataEnd.row - dataStart.row
        if dataAmount <= 1 :
            fontsize = 1400
            fontbold = True
        elif dataAmount <= 3 :
            fontsize = 1000
            fontbold = True
        elif dataAmount <= 5 :
            fontsize = 900
        elif dataAmount <= 7 :
            fontsize = 800

        labels_font = Font(typeface="Calibri")
        labels_character_properties = CharacterProperties(latin=labels_font, sz=fontsize, b=fontbold)
        labels_paragraph_properties = ParagraphProperties(defRPr=labels_character_properties)
        labels_paragraph = Paragraph(pPr=labels_paragraph_properties)
        labels_paragraphs = [labels_paragraph]
        labels_rich = RichText(p=labels_paragraphs)

        datalabel_list = DataLabelList(delete=False)
        datalabel_list.showSerName = False
        datalabel_list.showCatName = True
        datalabel_list.showVal = False
        datalabel_list.showPercent = True
        datalabel_list.showLegendKey = False
        datalabel_list.separator = "\n"
        datalabel_list.dLblPos = 'bestFit'
        datalabel_list.txPr = labels_rich
        self.chart.dataLabels = datalabel_list

        sheet.add_chart(self.chart, self.location)

        return
    
class CorrelationChart(Chart) :
    def __init__(self, stat:Stat, table:CorrelationTable, location:str) :
        super().__init__(stat=stat, table=table, location=location)
        self.targetCatArea = None
        self.comparisonCatArea = None

        startCell:Cell = self.sheet[self.areaTuple[0]]
        endCell:Cell = self.sheet[self.areaTuple[1]]

        targetCatStart = startCell.offset(1, 1)
        targetCatEnd = targetCatStart.offset(0, self.table.width-2)
        self.targetCatArea = targetCatStart.coordinate + ":" + targetCatEnd.coordinate

        comparisonCatStart = startCell.offset(2, 0)
        comparisonCatEnd = comparisonCatStart.offset(endCell.row - comparisonCatStart.row, 0)
        self.comparisonCatArea = comparisonCatStart.coordinate + ":" + comparisonCatEnd.coordinate

        dataStart = startCell.offset(2, 1)
        dataEnd:Cell = self.sheet[endCell.coordinate]
        self.dataArea = dataStart.coordinate + ":" + dataEnd.coordinate

        return
    
    def inflate(self, sh=None) :
        sheet:Worksheet = self.sheet
        if isinstance(sh, Worksheet) :
            sheet = sh
        else :
            super().inflate()

        # chart_space = ChartSpace(roundedCorners=False)
        # line_properties = LineProperties(round=False)
        # graphical_properties = GraphicalProperties(ln=line_properties)

        font = Font(typeface="Calibri")
        character_properties = CharacterProperties(latin=font, sz=1800, b=True)
        paragraph_properties = ParagraphProperties(defRPr=character_properties)

        self.chart = BarChart()
        self.chart.roundedCorners = False
        self.chart.x_axis.delete = False
        self.chart.y_axis.delete = False
        self.chart.title = self.table.title
        self.chart.title.tx.rich.p[0].pPr = paragraph_properties
        self.chart.title.overlay = False
        self.chart.type = "col"
        self.chart.grouping = "percentStacked"
        self.chart.overlap = 100
        self.chart.legend.overlay = False
        
        dataStart:Cell = self.sheet[self.targetCatArea][0][0]
        dataEnd:Cell = self.sheet[self.dataArea][-1][-1]

        catStart:Cell = self.sheet[self.comparisonCatArea][0][0]
        catEnd:Cell = self.sheet[self.comparisonCatArea][-1][-1]
        
        data = Reference(self.sheet, min_col=dataStart.column, min_row=dataStart.row, max_col=dataEnd.column, max_row=dataEnd.row)
        cats = Reference(self.sheet, min_col=catStart.column, min_row=catStart.row, max_row=catEnd.row)

        self.chart.add_data(data, titles_from_data=True)
        self.chart.set_categories(cats)
        self.chart.width = 12.8524
        self.chart.height = 7.7978

        # 열이 1개일 경우, 숨겨진 더미데이터 추가
        if len(self.chart.series) == 1 :
            dummydata = Reference(self.sheet, min_col=dataStart.column+1, min_row=dataStart.row, max_col=dataEnd.column+1, max_row=dataEnd.row)
            self.chart.add_data(dummydata, titles_from_data=True)
            self.chart.series[1].spPr.noFill = True

        sheet.add_chart(self.chart, self.location)

        return
